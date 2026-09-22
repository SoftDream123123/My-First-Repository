"""简历文件解析：把上传的 DOCX / PDF 抠成纯文本。

**纯函数模块**：不接 session、不落盘、不读配置。校验与算法都在这里，
错误一律 `ValidationError`（400）—— PRD §7 明确不新增状态码。

为什么宽 `except`：pypdf 与 python-docx 在损坏文件上抛出的异常谱系不稳定
（`PdfReadError` / `DependencyError` / `ValueError` / `xml` 层的各种异常），
逐个枚举必然漏。此处兜底是安全的，因为本模块**不落库、不落盘**：
宽捕获不会掩盖数据损坏，只会把「读不出来」如实告诉上传方。
"""

import io
import re
import zipfile

import pypdf
from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from app.core.exceptions import ValidationError

SUPPORTED_EXTENSIONS = (".docx", ".pdf")

# DOCX 解压后的展开体积上限。上传体积上限（配置项）拦不住 ZIP bomb ——
# 5 MB 的压缩包可以解出几个 GB；这一道单独拦。用中央目录的 file_size 求和，
# **不解压**，因此这道检查本身不吃内存。
DOCX_MAX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024

_MAGIC = {".pdf": b"%PDF-", ".docx": b"PK\x03\x04"}

_W_P = qn("w:p")
_W_TBL = qn("w:tbl")

# 连续 ≥ 2 个空行压成 1 个。PDF 抽取尤其容易产生成片空行（每行都是独立文本块）
_MANY_BLANK_LINES = re.compile(r"\n{3,}")

_TYPE_MESSAGE = "只支持 DOCX 或 PDF 文件，请换一个文件再试"
_PARSE_FAILED_MESSAGE = (
    "无法解析该文件，请确认它是完整的、未加密的 DOCX 或 PDF；"
    "也可以改为直接粘贴简历正文"
)
_ENCRYPTED_MESSAGE = (
    "该 PDF 已加密，请先解除密码保护后再上传；也可以改为直接粘贴简历正文"
)
_NO_TEXT_MESSAGE = (
    "未能从文件中提取到文字。它可能是扫描件或纯图片，"
    "请改为直接粘贴简历正文"
)


def parse_resume_file(data: bytes, filename: str | None) -> str:
    """把上传文件解析为纯文本。失败抛 `ValidationError`（400）。"""
    extension = _extension_of(filename)
    if extension not in SUPPORTED_EXTENSIONS:
        raise ValidationError(_TYPE_MESSAGE)

    # 扩展名由客户端提供，改个名字就能绕过 —— 按文件头魔数复核
    if not data.startswith(_MAGIC[extension]):
        raise ValidationError(_TYPE_MESSAGE)

    text = _parse_docx(data) if extension == ".docx" else _parse_pdf(data)
    normalized = _normalize(text)
    if not normalized:
        raise ValidationError(_NO_TEXT_MESSAGE)
    return normalized


def _extension_of(filename: str | None) -> str:
    """取最后一个点之后的部分并转小写；无扩展名或无文件名返回空串。

    `filename` 可能为 None（部分客户端不带文件名），此处必须容住，
    否则会抛 AttributeError 变成 500。
    """
    if not filename:
        return ""
    name = filename.strip()
    index = name.rfind(".")
    return name[index:].lower() if index >= 0 else ""


# ------------------------------- DOCX -------------------------------


def _parse_docx(data: bytes) -> str:
    try:
        _reject_zip_bomb(data)
        document = Document(io.BytesIO(data))
        lines: list[str] = []
        _collect_blocks(document.element.body, document, lines)
        return "\n".join(lines)
    except ValidationError:
        raise
    except Exception as exc:
        raise ValidationError(_PARSE_FAILED_MESSAGE) from exc


def _reject_zip_bomb(data: bytes) -> None:
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        total = sum(item.file_size for item in archive.infolist())
    if total > DOCX_MAX_UNCOMPRESSED_BYTES:
        raise ValidationError(_PARSE_FAILED_MESSAGE)


def _collect_blocks(parent, document, lines: list[str]) -> None:
    """按 body 的**子元素顺序**遍历，段落与表格保持原有版序。

    只取 `document.paragraphs` 会整块丢掉表格 —— 简历里的「教育经历」
    「工作经历」常做成表格，那样丢的是正文的一大块。
    """
    for child in parent.iterchildren():
        if child.tag == _W_P:
            lines.append(Paragraph(child, document).text)
        elif child.tag == _W_TBL:
            lines.extend(_table_lines(Table(child, document)))


def _table_lines(table) -> list[str]:
    lines: list[str] = []
    for row in table.rows:
        seen: set = set()
        cells: list[str] = []
        for cell in row.cells:
            # 合并单元格在 python-docx 里会让同一个 tc 出现在多个列位上，
            # 去重避免同一段文字被重复抽出来
            if cell._tc in seen:
                continue
            seen.add(cell._tc)
            cells.append(_cell_text(cell))
        line = "\t".join(cells).strip()
        if line:
            lines.append(line)
    return lines


def _cell_text(cell) -> str:
    """单元格文字：段落内换行压成空格，并**递归取嵌套表格**。

    嵌套表格与「漏掉表格」是同一类问题 —— 都会静默丢内容。
    """
    parts = [paragraph.text for paragraph in cell.paragraphs]
    for nested in cell.tables:
        parts.extend(_table_lines(nested))
    return " ".join(part for part in parts if part).strip()


# ------------------------------- PDF -------------------------------


def _parse_pdf(data: bytes) -> str:
    try:
        reader = pypdf.PdfReader(io.BytesIO(data))
        _ensure_readable(reader)
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    except ValidationError:
        raise
    except Exception as exc:
        raise ValidationError(_PARSE_FAILED_MESSAGE) from exc


def _ensure_readable(reader) -> None:
    """已加密则尝试用**空密码**解开，解不开才拒。

    很多导出工具默认给 PDF 加空密码，因此判定条件必须是
    「已加密 **且** 空密码解不开」—— 只看 `is_encrypted` 会误伤一批正常文件。
    """
    if not reader.is_encrypted:
        return
    try:
        unlocked = bool(reader.decrypt(""))
    except Exception:
        unlocked = False
    if not unlocked:
        raise ValidationError(_ENCRYPTED_MESSAGE)


# ----------------------------- 文本归一化 -----------------------------


def _normalize(text: str) -> str:
    """统一换行、去行尾空白、折叠连续空行、去首尾空白。

    **只做这四件事，不做重排**：不合并断行、不调顺序、不删页眉页脚残留。
    任何「更聪明」的处理都会让系统偷偷改掉候选人的原文，而他看到的正文
    正是要被 AI 消费的内容 —— 他必须看到真实结果。
    """
    unified = text.replace("\r\n", "\n").replace("\r", "\n")
    stripped = "\n".join(line.rstrip() for line in unified.split("\n"))
    return _MANY_BLANK_LINES.sub("\n\n", stripped).strip()
