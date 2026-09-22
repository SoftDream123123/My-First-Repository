"""resume 上传解析的单元测试。

**不经数据库、不经 HTTP** —— `parser.py` 是纯函数模块，`service` 那一层只多读
一个配置项，因此本文件直接把字节喂进去测。

测试用的 DOCX 与 PDF 都由 `tests/upload_helpers.py` 现场构造，**不需要往仓库里
放二进制测试资产**。
"""

import io

import pytest
from docx import Document

from app.core.config import get_settings
from app.core.exceptions import ValidationError
from app.resume import parser, service
from tests.upload_helpers import (
    blank_pdf_bytes,
    docx_bytes,
    encrypted_pdf_bytes,
    pdf_bytes,
)


def assert_400(callable_) -> ValidationError:
    """断言抛 400 `VALIDATION_ERROR` 并返回异常，供进一步断言文案。"""
    with pytest.raises(ValidationError) as caught:
        callable_()
    assert caught.value.status_code == 400
    assert caught.value.code == "VALIDATION_ERROR"
    return caught.value


# ------------------------------ DOCX ------------------------------


def test_docx_paragraphs_are_extracted_in_order():
    data = docx_bytes(["张三", "后端开发工程师", "五年经验"])
    text = parser.parse_resume_file(data, "resume.docx")
    assert text.split("\n") == ["张三", "后端开发工程师", "五年经验"]


def test_docx_table_text_is_extracted():
    """表格必须抠出来：简历里的教育/工作经历常做成表格。"""
    data = docx_bytes(
        [
            [["2019-2023", "某大学", "计算机"], ["2023-2026", "某公司", "后端"]],
        ]
    )
    text = parser.parse_resume_file(data, "resume.docx")
    assert "某大学" in text and "计算机" in text
    assert "某公司" in text and "后端" in text


def test_docx_paragraphs_and_tables_keep_document_order():
    data = docx_bytes(
        ["教育经历", [["某大学", "计算机"]], "工作经历", [["某公司", "后端"]]]
    )
    text = parser.parse_resume_file(data, "resume.docx")
    assert text.index("教育经历") < text.index("某大学")
    assert text.index("某大学") < text.index("工作经历")
    assert text.index("工作经历") < text.index("某公司")


def test_docx_merged_table_cell_is_not_duplicated():
    """合并单元格会让同一个 tc 出现在多个列位，去重后只应出现一次。"""
    document = Document()
    table = document.add_table(rows=1, cols=3)
    table.cell(0, 0).merge(table.cell(0, 2)).text = "合并的标题"
    buffer = io.BytesIO()
    document.save(buffer)

    text = parser.parse_resume_file(buffer.getvalue(), "resume.docx")
    assert text.count("合并的标题") == 1


def test_broken_docx_is_rejected_without_leaking_internals():
    error = assert_400(lambda: parser.parse_resume_file(b"PK\x03\x04" + b"junk", "a.docx"))
    assert "zip" not in error.message.lower()
    assert "lxml" not in error.message.lower()


def test_docx_uncompressed_size_limit_is_enforced(monkeypatch):
    """ZIP bomb 护栏：用 monkeypatch 把上限调到极小，避免造 50 MB 测试数据。"""
    monkeypatch.setattr(parser, "DOCX_MAX_UNCOMPRESSED_BYTES", 10)
    assert_400(lambda: parser.parse_resume_file(docx_bytes(["正文"]), "a.docx"))


# ------------------------------ PDF ------------------------------


def test_pdf_with_text_layer_is_extracted():
    text = parser.parse_resume_file(pdf_bytes("Python Backend Engineer"), "resume.pdf")
    assert "Python Backend Engineer" in text


def test_pdf_with_empty_password_still_parses():
    """很多导出工具默认加空密码，误伤这批文件是不可接受的。"""
    text = parser.parse_resume_file(encrypted_pdf_bytes(""), "resume.pdf")
    assert "Secret Resume" in text


def test_pdf_with_real_password_is_rejected():
    error = assert_400(lambda: parser.parse_resume_file(encrypted_pdf_bytes("hunter2"), "r.pdf"))
    assert "加密" in error.message


def test_pdf_without_text_layer_is_reported_as_scan():
    """扫描件没有「我是扫描件」的元数据，只能靠「取不到文字」倒推。"""
    error = assert_400(lambda: parser.parse_resume_file(blank_pdf_bytes(), "scan.pdf"))
    assert "扫描件" in error.message


def test_broken_pdf_is_rejected_without_leaking_internals():
    error = assert_400(lambda: parser.parse_resume_file(b"%PDF-1.4\njunk", "a.pdf"))
    assert "pypdf" not in error.message.lower()


# ------------------------------ 类型与魔数 ------------------------------


@pytest.mark.parametrize("filename", ["resume.doc", "resume.txt", "resume", "", None])
def test_unsupported_or_missing_filename_is_rejected(filename):
    """`filename` 可能为 None，必须容住，不能变成 AttributeError。"""
    assert_400(lambda: parser.parse_resume_file(pdf_bytes("x"), filename))


def test_extension_must_match_content():
    """扩展名由客户端提供，改个名字就能绕过 —— 魔数复核必须拦住。"""
    assert_400(lambda: parser.parse_resume_file(pdf_bytes("x"), "resume.docx"))
    assert_400(lambda: parser.parse_resume_file(docx_bytes(["x"]), "resume.pdf"))


# ------------------------------ 归一化 ------------------------------


def test_text_is_normalized():
    raw = docx_bytes(["第一行   ", "", "", "", "第二行"])
    text = parser.parse_resume_file(raw, "resume.docx")
    assert "\r" not in text
    assert "\n\n\n" not in text
    assert text == "第一行\n\n第二行"


# ------------------------------ service 层 ------------------------------


def test_empty_file_is_rejected_at_service_layer():
    error = assert_400(lambda: service.parse_uploaded_resume(b"", "resume.docx"))
    assert "空" in error.message


def test_oversize_file_is_rejected_with_limit_in_message():
    limit = get_settings().resume_upload_max_bytes
    data = b"PK\x03\x04" + b"0" * limit
    error = assert_400(lambda: service.parse_uploaded_resume(data, "resume.docx"))
    assert "5 MB" in error.message


def test_file_exactly_at_limit_passes_the_size_check():
    """边界：恰好等于上限**不算超**（超的是「大于」）。"""
    payload = docx_bytes(["正文"])
    limit = get_settings().resume_upload_max_bytes
    padded = payload + b"\x00" * (limit - len(payload))
    assert len(padded) == limit
    # 尾部补零会让 DOCX 结构损坏，因此这里只断言**不是**被体积校验拦下的
    with pytest.raises(ValidationError) as caught:
        service.parse_uploaded_resume(padded, "resume.docx")
    assert "MB" not in caught.value.message
