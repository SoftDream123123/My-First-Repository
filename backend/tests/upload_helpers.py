"""上传解析测试用的字节构造器（module-25）。

**不往仓库里放二进制测试资产**：
- DOCX 由 `python-docx` 现场生成；
- PDF 由 `pdf_bytes` 手写一个单页、带 `BT … Tj … ET` 内容流的最小文件
  （`pypdf` 能写结构但**排不出文字**，所以带文字层的 PDF 只能这样造）。

与 `sse_helpers.py` 同为跨测试文件共用的纯工具模块。
"""

import io

import pypdf
from docx import Document


def docx_bytes(blocks: list) -> bytes:
    """把 `blocks` 依次写进一个 DOCX。

    每个元素：`str` → 一个段落；`list[list[str]]` → 一张表格（外层是行）。
    """
    document = Document()
    for block in blocks:
        if isinstance(block, str):
            document.add_paragraph(block)
        else:
            table = document.add_table(rows=len(block), cols=len(block[0]))
            for row_index, row in enumerate(block):
                for col_index, value in enumerate(row):
                    table.cell(row_index, col_index).text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def pdf_bytes(text: str) -> bytes:
    """手写一个单页、带文字层的最小 PDF（约 600 字节）。"""
    content = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length "
        + str(len(content)).encode()
        + b" >>\nstream\n"
        + content
        + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]

    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for index, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(str(index).encode() + b" 0 obj\n" + body + b"\nendobj\n")

    xref = out.tell()
    out.write(b"xref\n0 " + str(len(objects) + 1).encode() + b"\n")
    out.write(b"0000000000 65535 f \n")
    for offset in offsets:
        out.write(f"{offset:010d} 00000 n \n".encode())
    out.write(
        b"trailer\n<< /Size "
        + str(len(objects) + 1).encode()
        + b" /Root 1 0 R >>\nstartxref\n"
        + str(xref).encode()
        + b"\n%%EOF\n"
    )
    return out.getvalue()


def blank_pdf_bytes() -> bytes:
    """只含一个空白页的 PDF —— 复现「扫描件 / 纯图片」：有页面但没有文字层。"""
    writer = pypdf.PdfWriter()
    writer.add_blank_page(width=612, height=792)
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def encrypted_pdf_bytes(password: str) -> bytes:
    writer = pypdf.PdfWriter(clone_from=io.BytesIO(pdf_bytes("Secret Resume")))
    writer.encrypt(password)
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()
